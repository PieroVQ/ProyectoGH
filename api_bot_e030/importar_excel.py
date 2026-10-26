"""Lee la plantilla vacía completada por el usuario y valida sin enviar datos."""
from __future__ import annotations
import argparse
import json
import math
from pathlib import Path
from openpyxl import load_workbook
from api_bot_e030.main import INPUTS, ValidationRequest, validate


def read_template(path: Path) -> dict:
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book["Entradas"]
        parameters, units = {}, {}
        for row in sheet.iter_rows(min_row=9, max_row=115, values_only=True):
            _, name, value, unit, *_ = row
            if name not in INPUTS or name in parameters:
                raise ValueError(f"Nombre de variable desconocido o repetido: {name}.")
            if value is not None and (isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value)):
                raise ValueError(f"{name}: ingrese un número finito, sin fórmula ni unidad en la celda.")
            parameters[name] = value
            units[name] = str(unit or "")
        if set(parameters) != set(INPUTS):
            raise ValueError("La plantilla no contiene las 107 variables exactas.")
        request = ValidationRequest(session_id="validacion-excel-local", parametros=parameters, unidades=units)
        _, errors, _ = validate(request)
        if errors:
            raise ValueError("\n".join(errors))
        return {"parametros": parameters, "unidades": units, "supuestos_aceptados": []}
    finally:
        book.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("excel", type=Path)
    parser.add_argument("--salida", type=Path, required=True)
    args = parser.parse_args()
    try:
        data = read_template(args.excel)
    except (ValueError, KeyError) as exc:
        parser.exit(2, f"No se creó la consulta:\n{exc}\n")
    args.salida.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"107 campos revisados; JSON creado localmente en {args.salida}. No se envió ningún dato.")
