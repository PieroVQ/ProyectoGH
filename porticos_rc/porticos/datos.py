"""Una hoja, una estructura por fila, sólo valores escalares.
openpyxl mantiene el proyecto Python portable, sin runtime JS propietario.
"""
from pathlib import Path
import json
import math
import numpy as np
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table,TableStyleInfo
from .normas import VERSION,R0,GAMMA_RC,LIM_DERIVA
from .geometria import GRUPOS

def scalar(v):
    if isinstance(v,np.generic): v=v.item()
    if isinstance(v,(list,dict,tuple,np.ndarray)): raise TypeError('Celda no escalar')
    if isinstance(v,float) and not math.isfinite(v): raise ValueError('Resultado no finito')
    return v

def a_fila(b,a,ident):
    m=a['metrado']; dr=a['deriva_normativa']; uu=a['u_normativa']
    row={'ID_estructura':ident,'ID_familia':b['familia'],'alternativa':b['variante'],'semilla_familia':b['semilla'],
         'pisos':b['n'],'vanos_X':len(b['lx']),'vanos_Y':len(b['ly']),
         'Lx_cm':sum(b['lx']),'Ly_cm':sum(b['ly']),'altura_total_cm':sum(b['alturas']),
         'fc_kgf_cm2':b['fc'],'fc_alcanza_21MPa':bool(b['fc']*.0980665>=21),'fy_kgf_cm2':b['fy'],'U':b['u'],'Z':b['z'],'perfil_suelo':b['suelo'],
         'S':b['s'],'TP_s':b['tp'],'TL_s':b['tl'],'codigo_cargas':b['codigo_cargas'],
         'qv_piso_kgf_cm2':b['qv'],'qv_cubierta_kgf_cm2':b['qv_techo'],
         'acabados_kgf_cm2':b['acabados'],'tabiques_fijos_kgf_cm2':b['tabiques'],
         'gamma_concreto_kgf_cm3':GAMMA_RC,'t_maciza_cm':b['t_maciza'],'t_aligerada_cm':b['t_aligerada'],
         'nervio_ancho_cm':10.,'nervio_separacion_cm':40.,'capa_compresion_cm':5.,
         'piso_inicio_bloque2':b['transicion']+1 if b['transicion'] else None,
         'cantidad_bloques':len(b['columnas']),'ratio_columna_muestreado':b['ratio_col'],
         'peso_gravitacional_kgf':m['peso_total_gravitacional'],'peso_sismico_kgf':float(sum(m['pesos'])),
         'fraccion_area_maciza':m['fraccion_maciza'],
         'Tx_s':a['Tx'],'Ty_s':a['Ty'],'T_estatico_s':a['T_estatico'],
         'R_asumido':R0,'regularidad_evaluable':a['regularidad_evaluable'],
         'regularidad_resistencia_verificada':False,'resistencia_elementos_verificada':False,
         'diafragma_rigido_asumido':True,'resistencia_diafragma_verificada':False,
         'segundo_orden_incluido':False,'normativa':VERSION,
         'criterio_columnas':'EMP_Pserv_0.35fc_MAS_GEOMETRIA_E060_21.6',
         'criterio_vigas':'EMP_L10_GRUPO_MAS_GEOMETRIA_E060_21.5',
         'criterio_losas':'E060_9.6.2_TAB9.1_9.6.3.3',
         'flecha_largo_plazo_verificada':False,'divisiones_desacopladas_flecha_asumidas':True,
         'distribucion_carga_vigas':'UNIFORME_EQUIVALENTE_RESULTANTE',
         'modelo_rigidez':'PORTICO3D_LINEAL_SECCION_BRUTA_SIN_BRAZOS_RIGIDOS',
         'error_equilibrio_relativo':a['error_equilibrio'],'error_reciprocidad':a['error_reciprocidad'],
         'error_modal_relativo':a['residuo_modal'],'iteraciones_predimensionamiento':b.get('iteraciones',0)}
    for d,l in [('X',b['lx']),('Y',b['ly'])]:
        row[f'luz_media_{d}_cm']=float(np.mean(l)); row[f'luz_min_{d}_cm']=min(l); row[f'luz_max_{d}_cm']=max(l)
        row[f'CV_luces_{d}']=float(np.std(l)/np.mean(l))
        for i in range(5): row[f'luz_{d}_{i+1}_cm']=l[i] if i<len(l) else None
    for bl in range(2):
        for g in GRUPOS:
            for p,dim in enumerate(('bX','bY')):
                row[f'col_{g}_bloque{bl+1}_{dim}_cm']=b['columnas'][bl][g][p] if bl<len(b['columnas']) else None
            row[f'col_{g}_bloque{bl+1}_area_requerida_cm2']=b['area_requerida'][bl][g] if bl<len(b['columnas']) else None
        for d in ('X','Y'):
            for p,dim in enumerate(('base','peralte')):
                row[f'viga_{d}_bloque{bl+1}_{dim}_cm']=b['vigas'][bl][d][p] if bl<len(b['vigas']) else None
    for i in range(5):
        for j in range(5):
            row[f'paño_X{i+1}_Y{j+1}_macizo']=bool(b['paños'][i][j]) if i<len(b['lx']) and j<len(b['ly']) else None
            for bl in range(2):
                active=i<len(b['lx']) and j<len(b['ly']) and bl<len(b['vigas'])
                direction=('X' if b['lx'][i]-b['vigas'][bl]['Y'][0]<=b['ly'][j]-b['vigas'][bl]['X'][0] else 'Y') if active else None
                row[f'paño_X{i+1}_Y{j+1}_direccion_corta_bloque{bl+1}']=direction
    for d,axis in enumerate(('X','Y')):
        row[f'V_estatico_{axis}_kgf']=a['V_estatico']
        row[f'V_modal_{axis}_kgf']=a['V_modal'][d,0]
        row[f'factor_escala_fuerzas_{axis}']=a['escala_fuerzas'][d]
        row[f'V_modal_ajustado_{axis}_kgf']=a['V_modal'][d,0]*a['escala_fuerzas'][d]
        row[f'V_combinado_{axis}_kgf']=a['V_combinado'][0,d]
        row[f'desplazamiento_techo_CM_{axis}_cm']=uu[-1,d,4]
        row[f'desplazamiento_techo_max_{axis}_cm']=max(uu[-1,d,:4])
        row[f'deriva_max_{axis}']=float(dr[:,d,:4].max())
        piso,corner=np.unravel_index(np.argmax(dr[:,d,:4]),(b['n'],4))
        row[f'piso_deriva_max_{axis}']=int(piso)+1
        row[f'esquina_deriva_max_{axis}']=int(corner)+1
        row[f'cumple_deriva_{axis}']=bool(dr[:,d,:4].max()<=LIM_DERIVA)
        row[f'razon_torsion_{axis}']=a['torsion'][d]
        row[f'indicador_PDelta_{axis}']=a['theta_indicador'][d]
        row[f'masa_modal_acumulada_{axis}']=sum(a['masa_modal'][d])
    for k in range(5):
        vals={'altura_cm':b['alturas'][k],'elevacion_cm':sum(b['alturas'][:k+1]),
              'peso_muerto_kgf':m['peso_muerto'][k],'peso_vivo_kgf':m['peso_vivo'][k],
              'peso_sismico_kgf':m['pesos'][k],'CM_X_cm':m['cm'][k,0],'CM_Y_cm':m['cm'][k,1],
              'inercia_peso_kgf_cm2':m['inercia_peso'][k]} if k<b['n'] else {name:None for name in
              ['altura_cm','elevacion_cm','peso_muerto_kgf','peso_vivo_kgf','peso_sismico_kgf','CM_X_cm','CM_Y_cm','inercia_peso_kgf_cm2']}
        for key,val in vals.items(): row[f'piso{k+1}_{key}']=val
        for d,axis in enumerate(('X','Y')):
            for name,arr in [('u_elastica_CM_cm',a['u_elastica']),('u_normativa_CM_cm',uu),('deriva_elastica_CM',a['deriva_elastica']),('deriva_normativa_CM',dr)]:
                row[f'piso{k+1}_{axis}_{name}']=arr[k,d,4] if k<b['n'] else None
            row[f'piso{k+1}_{axis}_rigidez_kgf_cm']=a['rigidez_pisos'][d,k] if k<b['n'] else None
            row[f'piso{k+1}_{axis}_deriva_max']=max(dr[k,d,:4]) if k<b['n'] else None
            row[f'piso{k+1}_{axis}_cortante_combinado_kgf']=a['V_combinado'][k,d] if k<b['n'] else None
            for p in range(4): row[f'piso{k+1}_{axis}_deriva_esquina{p+1}']=dr[k,d,p] if k<b['n'] else None
    for i in range(15):
        row[f'modo{i+1}_T_s']=a['periodos'][i] if i<len(a['periodos']) else None
        for d,axis in enumerate(('X','Y')):
            row[f'modo{i+1}_masa_{axis}']=a['masa_modal'][d,i] if i<len(a['periodos']) else None
    return {k:scalar(v) for k,v in row.items()}

def desde_fila(r):
    """Reconstrucción sin azar ni archivo auxiliar por edificio."""
    b={'familia':r['ID_familia'],'semilla':r['semilla_familia'],'variante':r['alternativa'],
       'n':r['pisos'],'lx':[r[f'luz_X_{i+1}_cm'] for i in range(r['vanos_X'])],
       'ly':[r[f'luz_Y_{i+1}_cm'] for i in range(r['vanos_Y'])],
       'alturas':[r[f'piso{k+1}_altura_cm'] for k in range(r['pisos'])],
       'fc':r['fc_kgf_cm2'],'fy':r['fy_kgf_cm2'],'u':r['U'],'z':r['Z'],'suelo':r['perfil_suelo'],
       's':r['S'],'tp':r['TP_s'],'tl':r['TL_s'],'codigo_cargas':r['codigo_cargas'],
       'qv':r['qv_piso_kgf_cm2'],'qv_techo':r['qv_cubierta_kgf_cm2'],
       'acabados':r['acabados_kgf_cm2'],'tabiques':r['tabiques_fijos_kgf_cm2'],
       'transicion':r['piso_inicio_bloque2']-1 if r['piso_inicio_bloque2'] else 0,
       'ratio_col':r['ratio_columna_muestreado'],'t_maciza':r['t_maciza_cm'],'t_aligerada':r['t_aligerada_cm'],
       'iteraciones':r['iteraciones_predimensionamiento'],'columnas':[],'vigas':[],'area_requerida':[]}
    for bl in range(r['cantidad_bloques']):
        b['columnas'].append({g:[r[f'col_{g}_bloque{bl+1}_{dim}_cm'] for dim in ('bX','bY')] for g in GRUPOS})
        b['vigas'].append({d:[r[f'viga_{d}_bloque{bl+1}_{dim}_cm'] for dim in ('base','peralte')] for d in ('X','Y')})
        b['area_requerida'].append({g:r[f'col_{g}_bloque{bl+1}_area_requerida_cm2'] for g in GRUPOS})
    b['paños']=[[int(r[f'paño_X{i+1}_Y{j+1}_macizo']) for j in range(r['vanos_Y'])] for i in range(r['vanos_X'])]
    return b

def exportar(rows,path):
    if not rows: return
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    wb=Workbook(); ws=wb.active; ws.title='Estructuras'
    headers=list(rows[0]); ws.append(headers)
    for row in rows: ws.append([scalar(row.get(k)) for k in headers])
    ws.freeze_panes='E2'; ws.row_dimensions[1].height=76
    for cell in ws[1]:
        cell.font=Font(name='Calibri',size=10,bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='17365D')
        cell.alignment=Alignment(wrap_text=True,vertical='center')
        ws.column_dimensions[cell.column_letter].width=24 if cell.column>4 else 16
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value,float): cell.number_format='0.000000' if ('deriva' in headers[cell.column-1] or 'qv_' in headers[cell.column-1]) else '0.0000'
            cell.font=Font(name='Calibri',size=10)
            if isinstance(cell.value,str) and len(cell.value)>24:
                cell.alignment=Alignment(wrap_text=True,vertical='center')
                ws.row_dimensions[cell.row].height=60
    tab=Table(displayName='BaseEstructuras',ref=f'A1:{get_column_letter(len(headers))}{len(rows)+1}')
    tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
    ws.add_table(tab)
    temporary=path.with_suffix('.tmp.xlsx'); wb.save(temporary); temporary.replace(path)

def leer_filas(path):
    wb=load_workbook(path,read_only=True,data_only=True)
    if len(wb.sheetnames)!=1: raise ValueError('Se esperaba exactamente una hoja')
    vals=wb.active.iter_rows(values_only=True); headers=next(vals)
    rows=[dict(zip(headers,row)) for row in vals if row[0] is not None]
    wb.close(); return rows

def diccionario(rows,path):
    if not rows: return
    labels={
      'ID_estructura':'Identificador único de fila; no predictor',
      'ID_familia':'Agrupar variantes para dividir entrenamiento y evaluación',
      'Tx_s':'Período del modo con mayor masa efectiva nominal en X',
      'Ty_s':'Período del modo con mayor masa efectiva nominal en Y',
      'R_asumido':'R=8 bajo hipótesis de ductilidad y regularidad; resistencia no verificada',
      'regularidad_evaluable':'Cumple criterios geométricos, masa, rigidez y torsión evaluados',
      'peso_gravitacional_kgf':'D+L completo, incluye peso directamente soportado por base',
      'peso_sismico_kgf':'Peso efectivo de niveles libres, D+fracción de L',
      'paño':'Patrón repetido en todos los pisos; booleano indica paño macizo',
    }
    result={}
    for k,v in rows[0].items():
        unit=next((u for u in ['kgf_cm3','kgf_cm2','kgf_cm','cm2','cm','kgf','s'] if k.endswith('_'+u)),'adimensional / categoría')
        unit={'kgf_cm3':'kgf/cm³','kgf_cm2':'kgf/cm²','kgf_cm':'kgf/cm','cm2':'cm²'}.get(unit,unit)
        if 'inercia_peso' in k: unit='kgf·cm²'
        output=any(k.startswith(p) for p in ['Tx','Ty','T_estatico','V_','deriva','desplazamiento','cumple','error','masa_modal','modo','regularidad','indicador','razon_torsion','factor_escala']) or (k.startswith('piso') and any(x in k for x in ['deriva','rigidez','u_elastica','u_normativa','cortante']))
        role='resultado' if output else 'entrada_o_descriptor'
        if k in ['ID_estructura','ID_familia','alternativa','semilla_familia','normativa'] or 'criterio' in k or 'verificad' in k or 'asumid' in k or k=='fc_alcanza_21MPa': role='metadato_no_predictor'
        result[k]={'descripcion':labels.get(k,k.replace('_',' ')), 'unidad':unit,'rol':role,'ausente':'celda vacía: no aplicable'}
    Path(path).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    features=['fc_kgf_cm2','fy_kgf_cm2','pisos','vanos_X','vanos_Y','Lx_cm','Ly_cm','altura_total_cm','piso1_altura_cm','piso2_altura_cm',
              'luz_media_X_cm','luz_media_Y_cm','CV_luces_X','CV_luces_Y','U','Z','S','TP_s','TL_s',
              'qv_piso_kgf_cm2','acabados_kgf_cm2','tabiques_fijos_kgf_cm2','peso_sismico_kgf',
              't_maciza_cm','t_aligerada_cm','fraccion_area_maciza','cantidad_bloques','piso_inicio_bloque2']
    features += [k for k in rows[0] if (k.startswith('col_') or k.startswith('viga_')) and k.endswith('_cm')]
    manifest={'entradas_iniciales':features,'grupo':'ID_familia','objetivos':['Tx_s','Ty_s','V_modal_ajustado_X_kgf','V_modal_ajustado_Y_kgf','desplazamiento_techo_max_X_cm','desplazamiento_techo_max_Y_cm','deriva_max_X','deriva_max_Y'],
              'nota':'Validación por familias. Ausencias estructurales: usar indicador y tratamiento aprendido sólo en entrenamiento. Revisar correlaciones y variables constantes. No se entrena ML.'}
    Path(path).with_name('variables_ml.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
